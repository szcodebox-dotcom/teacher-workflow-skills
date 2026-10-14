"""Artifact-level checks for the self-contained teaching deck generator."""

import copy
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "teacher-slide-deck"
SPEC = importlib.util.spec_from_file_location("teacher_deck_builder", SKILL / "scripts" / "build_deck.py")
BUILDER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BUILDER)


class DeckArtifactTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = json.loads((SKILL / "references" / "example-deck.json").read_text(encoding="utf-8"))

    def test_output_has_all_pages_answers_notes_and_editable_table(self):
        with tempfile.TemporaryDirectory() as directory:
            source, target = Path(directory) / "source.json", Path(directory) / "lesson.pptx"
            source.write_text(json.dumps(self.data, ensure_ascii=False), encoding="utf-8")
            self.assertEqual(BUILDER.build_deck(source, target), 8)
            deck = Presentation(target)
            self.assertAlmostEqual(deck.slide_width / deck.slide_height, 16 / 9, places=5)
            self.assertEqual(len(deck.slides), 8)
            for slide, page in zip(deck.slides, self.data["slides"]):
                texts = [shape.text for shape in slide.shapes if shape.has_text_frame]
                self.assertIn(page["title"], texts)
                self.assertIn(page["notes"], slide.notes_slide.notes_text_frame.text)
            tables = [shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table]
            self.assertEqual(len(tables), 1)
            self.assertEqual(tables[0].cell(1, 4).text, "10 枚")
            self.assertIn("28÷4=7", deck.slides[6].notes_slide.notes_text_frame.text)
            self.assertNotIn("28÷4=7", "\n".join(shape.text for shape in deck.slides[6].shapes if shape.has_text_frame))
            # Verify actual edit/save/reopen rather than merely looking for XML.
            tables[0].cell(1, 4).text = "11 枚"
            title_shape = next(shape for shape in deck.slides[0].shapes if shape.has_text_frame and shape.text == self.data["slides"][0]["title"])
            title_shape.text = "可编辑标题已修改"
            edited = Path(directory) / "edited.pptx"
            deck.save(edited)
            reopened = Presentation(edited)
            self.assertTrue(any(shape.has_text_frame and shape.text == "可编辑标题已修改" for shape in reopened.slides[0].shapes))
            self.assertEqual(next(shape.table for shape in reopened.slides[1].shapes if shape.has_table).cell(1, 4).text, "11 枚")

    def test_missing_teaching_notes_is_rejected(self):
        data = copy.deepcopy(self.data)
        data["slides"][0]["notes"] = " "
        with self.assertRaisesRegex(ValueError, "notes"):
            BUILDER.validate_deck(data)

    def test_ragged_data_table_is_rejected(self):
        data = copy.deepcopy(self.data)
        data["slides"][1]["rows"][0].pop()
        with self.assertRaisesRegex(ValueError, "header width"):
            BUILDER.validate_deck(data)

    def test_class_time_mismatch_is_rejected(self):
        data = copy.deepcopy(self.data)
        data["total_minutes"] = 45
        with self.assertRaisesRegex(ValueError, "minutes total"):
            BUILDER.validate_deck(data)

    def test_invalid_input_keeps_existing_artifact(self):
        with tempfile.TemporaryDirectory() as directory:
            source, target = Path(directory) / "bad.json", Path(directory) / "existing.pptx"
            target.write_bytes(b"existing artifact remains intact")
            data = copy.deepcopy(self.data)
            data["slides"][0]["minutes"] = -1
            source.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            with self.assertRaises(ValueError):
                BUILDER.build_deck(source, target)
            self.assertEqual(target.read_bytes(), b"existing artifact remains intact")

    def test_long_panel_is_rejected_instead_of_silent_overflow(self):
        data = copy.deepcopy(self.data)
        data["slides"][2]["panels"][0]["lines"] = ["很" * 50] * 3
        with self.assertRaisesRegex(ValueError, "140"):
            BUILDER.validate_deck(data)

    def test_wrapped_cjk_lines_are_rejected_when_below_character_limit(self):
        data = copy.deepcopy(self.data)
        # 130 characters passes the coarse 140-character bound, but would
        # wrap into 9 visual lines in the narrow panel plus paragraph gaps.
        data["slides"][2]["panels"][0]["lines"] = ["长" * 45, "中" * 45, "文" * 40]
        with self.assertRaisesRegex(ValueError, "may overflow"):
            BUILDER.validate_deck(data)

    def test_dense_table_cell_is_rejected_for_available_row_height(self):
        data = copy.deepcopy(self.data)
        data["slides"][1]["rows"] = [["完整教学记录内容"] * 5 for _ in range(6)]
        with self.assertRaisesRegex(ValueError, "may overflow"):
            BUILDER.validate_deck(data)

    def test_non_finite_class_time_is_rejected(self):
        for value in [float("nan"), float("inf"), float("-inf")]:
            data = copy.deepcopy(self.data)
            data["total_minutes"] = value
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "positive number"):
                BUILDER.validate_deck(data)

    def test_heading_that_wraps_out_of_its_box_is_rejected(self):
        data = copy.deepcopy(self.data)
        data["slides"][2]["panels"][0]["heading"] = "这段标题虽然字符不多但在双栏中需要换行"
        with self.assertRaisesRegex(ValueError, "panel heading.*overflow"):
            BUILDER.validate_deck(data)


if __name__ == "__main__":
    unittest.main()
