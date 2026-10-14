#!/usr/bin/env python3
"""Run all four bundled generators from a relocated, extracted expert package."""
import argparse
import json
from pathlib import Path
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[1]


def run(output, skip_pptx=False):
    output = Path(output).resolve()
    if output.exists() and any(output.iterdir()):
        raise ValueError('Output directory must be empty: ' + str(output))
    if not skip_pptx:
        try:
            import pptx  # noqa: F401
        except ImportError as exc:
            raise ValueError('Install requirements-demo.txt with this Python before generating PPTX') from exc
    output.mkdir(parents=True, exist_ok=True)
    results = []
    for name in ['teacher-exam-builder', 'teacher-blueprint-exam']:
        folder = ROOT / 'skills' / name
        proc = subprocess.run([sys.executable, str(folder/'scripts/audit_exam.py'), str(folder/'references/example-exam.json')], capture_output=True, text=True, check=True)
        data = json.loads(proc.stdout)
        (output/(name+'-audit.json')).write_text(json.dumps(data, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
        results.append({'skill':name, 'executed':True, 'audit':data})
    folder = ROOT/'skills/teacher-learning-insights'
    subprocess.run([sys.executable,str(folder/'scripts/analyze_scores.py'),str(folder/'examples/scores.csv'),'--output-dir',str(output/'learning')],check=True,capture_output=True,text=True)
    assert (output/'learning/report.json').is_file() and (output/'learning/report.md').is_file()
    results.append({'skill':'teacher-learning-insights','executed':True,'files':['learning/report.json','learning/report.md']})
    if not skip_pptx:
        folder=ROOT/'skills/teacher-slide-deck'
        subprocess.run([sys.executable,str(folder/'scripts/build_deck.py'),str(folder/'references/example-deck.json'),'--output',str(output/'lesson.pptx')],check=True,capture_output=True,text=True)
        from pptx import Presentation
        deck=Presentation(output/'lesson.pptx')
        assert len(deck.slides)==8
        assert all(slide.notes_slide.notes_text_frame.text.strip() for slide in deck.slides)
        assert sum(shape.has_table for slide in deck.slides for shape in slide.shapes)==1
        results.append({'skill':'teacher-slide-deck','executed':True,'file':'lesson.pptx','slides':8,'notes':8,'editable_tables':1,'visual_rendering_checked':False})
    else:
        results.append({'skill':'teacher-slide-deck','executed':False,'reason':'explicit --skip-pptx'})
    report={'synthetic_examples':True,'results':results,'output_dir':str(output)}
    (output/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return report


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',type=Path,default=Path.cwd()/'teacher-expert-demo')
    parser.add_argument('--skip-pptx',action='store_true',help='Run only the three standard-library scripts')
    args=parser.parse_args()
    try:
        print(json.dumps(run(args.output_dir,args.skip_pptx),ensure_ascii=False,indent=2))
    except (ValueError,AssertionError,OSError,subprocess.SubprocessError) as exc:
        parser.exit(1,str(exc)+'\n')


if __name__=='__main__':
    main()
