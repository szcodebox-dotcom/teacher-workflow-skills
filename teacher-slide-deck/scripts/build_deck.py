#!/usr/bin/env python3
"""Build an editable 16:9 teaching deck from complete, validated JSON.

Runtime: Python 3.10+ and python-pptx 1.x. No bundled theme, root-level
configuration, browser, external media or other skill is required.
"""

from __future__ import annotations

import argparse
import json
import math
import os
from pathlib import Path
import tempfile
import unicodedata

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


NAVY = "14354E"
TEAL = "087F8C"
PAPER = "F4F8FA"
WHITE = "FFFFFF"
INK = "243746"
MUTED = "516A7C"
PALE = "E5F2F3"


def require_text(value, label, max_length=120):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{label} must be non-empty text")
    if len(value) > max_length:
        raise ValueError(f"{label} exceeds {max_length} characters; split the content")
    return value


def estimated_lines(text, width_points, font_points):
    """Conservative wrapping estimate; CJK takes one em, Latin 0.65 em.

    This is a preflight bound, not a claim to reproduce a rendering engine.
    Explicit newlines start fresh visual lines. A safety inset limits width.
    """
    capacity = max(1, (width_points - 8) / font_points)
    total = 0
    for line in text.split("\n"):
        used, count = 0.0, 1
        for character in line:
            advance = 1.0 if unicodedata.east_asian_width(character) in {"W", "F"} else 0.65
            if used + advance > capacity:
                count += 1
                used = 0.0
            used += advance
        total += count
    return total


def reject_if_overflow(lines, width, height, font_size, paragraph_gap, label):
    visual_lines = sum(estimated_lines(line, width, font_size) for line in lines)
    required = visual_lines * font_size * 1.2 + max(0, len(lines) - 1) * paragraph_gap
    if required > height - 4:
        raise ValueError(f"{label}: content may overflow the layout ({visual_lines} estimated lines); split the slide")


def validate_deck(data):
    """Validate teaching essentials and conservative bounds for the layouts."""
    if not isinstance(data, dict):
        raise ValueError("deck must be a JSON object")
    require_text(data.get("title"), "title", 80)
    require_text(data.get("audience"), "audience", 70)
    require_text(data.get("font", "Microsoft YaHei"), "font", 80)
    slides = data.get("slides")
    if not isinstance(slides, list) or not 1 <= len(slides) <= 60:
        raise ValueError("slides must contain 1–60 slides")
    total_minutes = 0
    for number, slide in enumerate(slides, 1):
        label = f"slide {number}"
        if not isinstance(slide, dict):
            raise ValueError(f"{label} must be an object")
        kind = slide.get("kind")
        if kind not in {"title", "content", "table"}:
            raise ValueError(f"{label}: kind must be title, content or table")
        require_text(slide.get("title"), f"{label} title", 34)
        require_text(slide.get("notes"), f"{label} notes", 6000)
        cover = kind == "title"
        reject_if_overflow([slide["title"]], 11.66 * 72, (1.31 if cover else 0.64) * 72, 40 if cover else 32, 0, f"{label} title")
        reject_if_overflow([data["audience"]], 11.66 * 72, (0.41 if cover else 0.26) * 72, 16 if cover else 12, 0, f"{label} audience")
        minutes = slide.get("minutes")
        if isinstance(minutes, bool) or not isinstance(minutes, (int, float)) or not math.isfinite(minutes) or not 0 < minutes <= 120:
            raise ValueError(f"{label}: minutes must be a positive number up to 120")
        total_minutes += minutes
        if "takeaway" in slide:
            require_text(slide["takeaway"], f"{label} takeaway", 54)
            reject_if_overflow([slide["takeaway"]], 11.66 * 72, 0.61 * 72, 17, 0, f"{label} takeaway")
        if kind == "title":
            require_text(slide.get("subtitle"), f"{label} subtitle", 100)
            reject_if_overflow([slide["subtitle"]], 10.96 * 72, 1.51 * 72, 25, 0, f"{label} subtitle")
        elif kind == "content":
            panels = slide.get("panels")
            if not isinstance(panels, list) or not 1 <= len(panels) <= 2:
                raise ValueError(f"{label}: content needs one or two panels")
            for index, panel in enumerate(panels, 1):
                if not isinstance(panel, dict):
                    raise ValueError(f"{label} panel {index} must be an object")
                require_text(panel.get("heading"), f"{label} panel heading", 24)
                lines = panel.get("lines")
                if not isinstance(lines, list) or not 1 <= len(lines) <= 5:
                    raise ValueError(f"{label}: panel needs 1–5 lines")
                for line in lines:
                    require_text(line, f"{label} panel line", 64)
                if sum(len(line) for line in lines) > 140:
                    raise ValueError(f"{label}: panel body exceeds 140 characters; split the slide")
                panel_width = (11.7 - 0.35 * (len(panels) - 1)) / len(panels)
                reject_if_overflow([panel["heading"]], (panel_width - 0.74) * 72, 0.51 * 72, 24, 0, f"{label} panel heading")
                reject_if_overflow(lines, (panel_width - 0.74) * 72, 2.88 * 72, 23, 12, f"{label} panel {index}")
        else:
            headers, rows = slide.get("headers"), slide.get("rows")
            if not isinstance(headers, list) or not 2 <= len(headers) <= 5:
                raise ValueError(f"{label}: table needs 2–5 columns")
            for header in headers:
                require_text(header, f"{label} table header", 14)
            if not isinstance(rows, list) or not 1 <= len(rows) <= 6:
                raise ValueError(f"{label}: table needs 1–6 data rows")
            for row in rows:
                if not isinstance(row, list) or len(row) != len(headers):
                    raise ValueError(f"{label}: each row must match the header width")
                for cell in row:
                    require_text(cell, f"{label} table cell", 24)
            column_width = (11.7 / len(headers) - 0.24) * 72
            row_height = (3.9 / (len(rows) + 1) - 0.12) * 72
            for cell in headers:
                reject_if_overflow([cell], column_width, row_height, 20, 0, f"{label} table header")
            for row in rows:
                for cell in row:
                    reject_if_overflow([cell], column_width, row_height, 22, 0, f"{label} table cell")
    if "total_minutes" in data:
        expected = data["total_minutes"]
        if isinstance(expected, bool) or not isinstance(expected, (int, float)) or not math.isfinite(expected) or expected <= 0:
            raise ValueError("total_minutes must be a positive number")
        if abs(total_minutes - expected) > 0.01:
            raise ValueError(f"slide minutes total {total_minutes:g}, expected {expected:g}")
    return data


def rgb(value):
    return RGBColor.from_string(value)


def apply_font(paragraph, name, size, color, bold=False):
    font = paragraph.font
    font.name = name
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = rgb(color)
    # Set the East Asian face as well as the Latin face; Chinese otherwise
    # inherits a theme font in some presentation applications.
    for tag in ("a:ea", "a:cs"):
        element = font._element.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            font._element.append(element)
        element.set("typeface", name)


def rectangle(slide, x, y, w, h, color, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def text_box(slide, x, y, w, h, lines, font, size=24, color=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        apply_font(paragraph, font, size, color, bold)
        paragraph.space_after = Pt(12)
        paragraph.line_spacing = 1.15
    return box


def write_table(slide, spec, font):
    headers, rows = spec["headers"], spec["rows"]
    row_count = len(rows) + 1
    table = slide.shapes.add_table(row_count, len(headers), Inches(0.8), Inches(1.8), Inches(11.7), Inches(3.9)).table
    for index in range(row_count):
        table.rows[index].height = Inches(3.9 / row_count)
        values = headers if index == 0 else rows[index - 1]
        for column, value in enumerate(values):
            cell = table.cell(index, column)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(TEAL if index == 0 else (WHITE if index % 2 else PALE))
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text = value
            for paragraph in cell.text_frame.paragraphs:
                apply_font(paragraph, font, 20 if index == 0 else 22, WHITE if index == 0 else INK, index == 0)
    return table


def create_presentation(data):
    validate_deck(data)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = data["title"]
    prs.core_properties.subject = data["audience"]
    prs.core_properties.author = "Teacher Workflow Pack"
    font = data.get("font", "Microsoft YaHei")
    count = len(data["slides"])
    for number, spec in enumerate(data["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        cover = spec["kind"] == "title"
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(NAVY if cover else PAPER)
        rectangle(slide, 0, 0, 0.16, 7.5, TEAL)
        if cover:
            rectangle(slide, 0.8, 1.45, 1.0, 0.09, "5AC5C8")
            text_box(slide, 0.8, 0.55, 11.7, 0.45, [data["audience"]], font, 16, "ADD9E0")
            text_box(slide, 0.8, 2.0, 11.7, 1.35, [spec["title"]], font, 40, WHITE, True)
            text_box(slide, 0.8, 3.75, 11.0, 1.55, [spec["subtitle"]], font, 25, WHITE)
        else:
            text_box(slide, 0.8, 0.35, 11.7, 0.3, [data["audience"]], font, 12, MUTED)
            text_box(slide, 0.8, 0.85, 11.7, 0.68, [spec["title"]], font, 32, NAVY, True)
            if spec["kind"] == "table":
                write_table(slide, spec, font)
            else:
                panels = spec["panels"]
                gap = 0.35
                width = (11.7 - gap * (len(panels) - 1)) / len(panels)
                for index, panel in enumerate(panels):
                    x = 0.8 + index * (width + gap)
                    rectangle(slide, x, 1.8, width, 4.25, WHITE, True)
                    rectangle(slide, x + 0.25, 2.07, 0.055, 0.45, TEAL)
                    text_box(slide, x + 0.45, 2.02, width - 0.7, 0.55, [panel["heading"]], font, 24, TEAL, True)
                    text_box(slide, x + 0.35, 2.9, width - 0.7, 2.92, panel["lines"], font, 23)
        if spec.get("takeaway"):
            text_box(slide, 0.8, 6.3, 11.7, 0.65, [spec["takeaway"]], font, 17, WHITE if cover else TEAL)
        text_box(slide, 0.8, 7.05, 10.0, 0.28, [data["title"]], font, 10, "ADD9E0" if cover else MUTED)
        text_box(slide, 11.7, 7.05, 0.8, 0.28, [f"{number:02d} / {count:02d}"], font, 10, "ADD9E0" if cover else MUTED)
        slide.notes_slide.notes_text_frame.text = f"建议用时：{spec['minutes']:g} 分钟\n{spec['notes']}"
    return prs


def build_deck(input_path, output_path):
    source, target = Path(input_path), Path(output_path)
    if target.suffix.lower() != ".pptx":
        raise ValueError("output must have a .pptx extension")
    if source.resolve() == target.resolve():
        raise ValueError("input and output must be different files")
    with source.open("r", encoding="utf-8-sig") as handle:
        data = json.load(handle)
    prs = create_presentation(data)
    target.parent.mkdir(parents=True, exist_ok=True)
    # Save atomically, so a failed write cannot leave a partial presentation.
    with tempfile.NamedTemporaryFile(suffix=".pptx", prefix=".teacher-deck-", dir=target.parent, delete=False) as handle:
        temp_path = Path(handle.name)
    try:
        prs.save(temp_path)
        reopened = Presentation(temp_path)
        if len(reopened.slides) != len(data["slides"]):
            raise RuntimeError("written presentation has an unexpected slide count")
        for number, slide in enumerate(reopened.slides, 1):
            if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip():
                raise RuntimeError(f"written slide {number} has no teaching notes")
        os.replace(temp_path, target)
    finally:
        temp_path.unlink(missing_ok=True)
    return len(prs.slides)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Complete UTF-8 deck JSON")
    parser.add_argument("--output", required=True, type=Path, help="Destination .pptx")
    args = parser.parse_args(argv)
    try:
        count = build_deck(args.input, args.output)
    except (OSError, ValueError, TypeError) as exc:
        parser.exit(2, f"Cannot build presentation: {exc}\n")
    print(f"Created {args.output.resolve()} ({count} editable slides with teaching notes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
